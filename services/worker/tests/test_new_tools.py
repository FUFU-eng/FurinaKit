"""新增的 5 个能力的测试（video-trim / audio-trim / video-frame-extract / csv-excel / markdown-to-pdf）。

- 不联网：视频 / 音频素材全部用 `ffmpeg -f lavfi` 现场生成
- venv 里没有 pytest，用 stdlib unittest：`python -m unittest discover -s tests -v`
- 既有工具（images-to-pdf / pdf-merge / video-compress / audio-merge）的回归测试见
  tests/test_existing_tools_regression.py
"""

import os
import re
import unittest
from pathlib import Path
from unittest import mock

import pymupdf as fitz
from openpyxl import Workbook, load_workbook
from PIL import Image

from app.tools import audio_trim, csv_excel, markdown_pdf, video_trim
from tests._job_harness import (
    WorkerToolTest,
    ffmpeg_path,
    make_audio,
    make_video,
    probe_duration,
    run_tool_cmd,
    stream_types,
)


# ══════════════════════════════════════════════════════════════════════
# 1) video-trim
# ══════════════════════════════════════════════════════════════════════
class VideoTrimTests(WorkerToolTest):
    def setUp(self):
        super().setUp()
        self.video = make_video(self.work / "clip.mp4", duration=5, gop=10)

    def out(self, name="out.mp4"):
        return str(self.work / name)

    def test_fast_copy_trim_is_playable(self):
        out = self.out("fast.mp4")
        result = video_trim.trim_video(self.video, out, start=1, end=3, mode="fast")
        self.assertTrue(result["success"], result.get("error"))
        self.assertTrue(os.path.isfile(out))
        self.assertGreater(os.path.getsize(out), 1024)
        self.assertIn("video", stream_types(out))
        duration = probe_duration(out)
        self.assertGreater(duration, 1.0, "裁剪后的时长不合理")
        self.assertLess(duration, 3.8)
        self.assertIn("裁剪完成", result["message"])

    def test_precise_trim_duration_is_exact(self):
        out = self.out("precise.mp4")
        result = video_trim.trim_video(self.video, out, start="00:00:01", end="00:00:03", mode="precise")
        self.assertTrue(result["success"], result.get("error"))
        self.assertAlmostEqual(probe_duration(out), 2.0, delta=0.35)
        self.assertIn("精确重编码", result["message"])

    def test_default_end_trims_to_tail(self):
        out = self.out("tail.mp4")
        result = video_trim.trim_video(self.video, out, start="00:00:03", end=None, mode="precise")
        self.assertTrue(result["success"], result.get("error"))
        self.assertAlmostEqual(probe_duration(out), 2.0, delta=0.35)

    def test_default_start_is_zero(self):
        out = self.out("head.mp4")
        result = video_trim.trim_video(self.video, out, start=None, end="2", mode="precise")
        self.assertTrue(result["success"], result.get("error"))
        self.assertAlmostEqual(probe_duration(out), 2.0, delta=0.35)

    def test_hhmmss_with_millis(self):
        out = self.out("ms.mp4")
        result = video_trim.trim_video(self.video, out, start="00:00:01.500", end="00:00:02.500",
                                      mode="precise")
        self.assertTrue(result["success"], result.get("error"))
        self.assertAlmostEqual(probe_duration(out), 1.0, delta=0.3)

    def test_end_before_start_reports_human_error(self):
        out = self.out("bad.mp4")
        result = video_trim.trim_video(self.video, out, start=4, end=2, mode="fast")
        self.assertFalse(result["success"])
        self.assertIn("结束时间必须晚于开始时间", result["error"])
        self.assertFalse(os.path.exists(out), "失败时不应该留下半成品")

    def test_invalid_time_format_reports_human_error(self):
        for bad in ("abc", "12:99", "1:2:3:4", "::", "", " 1h "):
            if bad == "":
                continue  # 空字符串按「未提供」处理，不算非法
            with self.subTest(bad=bad):
                result = video_trim.trim_video(self.video, self.out("bad2.mp4"), start=bad)
                self.assertFalse(result["success"], f"{bad!r} 应该被拒绝")
                self.assertIn("格式不正确", result["error"])

    def test_negative_time_reports_human_error(self):
        result = video_trim.trim_video(self.video, self.out("neg.mp4"), start=-5)
        self.assertFalse(result["success"])
        self.assertIn("不能是负数", result["error"])

    def test_start_beyond_duration_reports_duration(self):
        result = video_trim.trim_video(self.video, self.out("far.mp4"), start=30, end=None)
        self.assertFalse(result["success"])
        self.assertIn("这个视频只有", result["error"])
        self.assertIn("超出视频长度", result["error"])

    def test_end_beyond_duration_is_clamped(self):
        out = self.out("clamp.mp4")
        result = video_trim.trim_video(self.video, out, start=4, end=99, mode="precise")
        self.assertTrue(result["success"], result.get("error"))
        self.assertAlmostEqual(probe_duration(out), 1.0, delta=0.4)
        self.assertIn("已自动裁到结尾", result["message"])

    def test_missing_file(self):
        result = video_trim.trim_video(str(self.work / "nope.mp4"), self.out("x.mp4"))
        self.assertFalse(result["success"])
        self.assertIn("找不到", result["error"])

    def test_audio_file_is_rejected(self):
        audio = make_audio(self.work / "tone.mp3", duration=2)
        result = video_trim.trim_video(audio, self.out("fromaudio.mp4"))
        self.assertFalse(result["success"])
        self.assertIn("没有视频画面", result["error"])

    def test_bad_mode(self):
        result = video_trim.trim_video(self.video, self.out("m.mp4"), start=0, end=1, mode="turbo")
        self.assertFalse(result["success"])
        self.assertIn("裁剪模式只支持", result["error"])

    def test_falls_back_to_reencode_when_copy_fails(self):
        """流复制失败时必须自动改用重编码，并在消息里说明。"""
        out = self.out("fallback.mp4")
        real_run = video_trim.run_process

        class FakeResult:
            returncode = 1
            stdout = ""
            stderr = "simulated copy failure"

        def fake_run(cmd, timeout=None):
            if "copy" in cmd:
                return FakeResult()
            return real_run(cmd, timeout=timeout)

        with mock.patch.object(video_trim, "run_process", fake_run):
            result = video_trim.trim_video(self.video, out, start=1, end=2, mode="fast")

        self.assertTrue(result["success"], result.get("error"))
        self.assertIn("自动改用精确重编码", result["message"])
        self.assertTrue(os.path.isfile(out))


# ══════════════════════════════════════════════════════════════════════
# 2) audio-trim
# ══════════════════════════════════════════════════════════════════════
class AudioTrimTests(WorkerToolTest):
    def setUp(self):
        super().setUp()
        self.mp3 = make_audio(self.work / "tone.mp3", duration=4)

    def out(self, name):
        return str(self.work / name)

    def test_fast_trim_keeps_extension_and_duration(self):
        out = self.out("tone-trim.mp3")
        result = audio_trim.trim_audio(self.mp3, out, start=1, end=3, mode="fast")
        self.assertTrue(result["success"], result.get("error"))
        self.assertTrue(out.endswith(".mp3"))
        self.assertEqual(result["mime"], "audio/mpeg")
        self.assertAlmostEqual(probe_duration(out), 2.0, delta=0.5)
        self.assertIn("裁剪完成", result["message"])

    def test_precise_trim_duration(self):
        out = self.out("tone-precise.mp3")
        result = audio_trim.trim_audio(self.mp3, out, start="00:00:01", end="00:00:03", mode="precise")
        self.assertTrue(result["success"], result.get("error"))
        self.assertAlmostEqual(probe_duration(out), 2.0, delta=0.3)
        self.assertIn("重编码", result["message"])

    def test_default_end_trims_to_tail(self):
        out = self.out("mp3-tail.mp3")
        result = audio_trim.trim_audio(self.mp3, out, start=3, end=None, mode="precise")
        self.assertTrue(result["success"], result.get("error"))
        self.assertAlmostEqual(probe_duration(out), 1.0, delta=0.35)

    def test_wav_container_keeps_extension(self):
        wav = make_audio(self.work / "tone.wav", duration=3, fmt="wav")
        out = self.out("tone-trim.wav")
        result = audio_trim.trim_audio(wav, out, start=1, end=2, mode="fast")
        self.assertTrue(result["success"], result.get("error"))
        self.assertTrue(out.endswith(".wav"))
        self.assertEqual(result["mime"], "audio/wav")
        self.assertAlmostEqual(probe_duration(out), 1.0, delta=0.35)

    def test_flac_container_keeps_extension(self):
        flac = make_audio(self.work / "tone.flac", duration=3, fmt="flac")
        out = self.out("tone-trim.flac")
        result = audio_trim.trim_audio(flac, out, start=1, end=2, mode="fast")
        self.assertTrue(result["success"], result.get("error"))
        self.assertTrue(out.endswith(".flac"))
        self.assertAlmostEqual(probe_duration(out), 1.0, delta=0.35)

    def test_ogg_container_keeps_extension(self):
        ogg = make_audio(self.work / "tone.ogg", duration=3, fmt="ogg")
        out = self.out("tone-trim.ogg")
        result = audio_trim.trim_audio(ogg, out, start=1, end=2, mode="fast")
        self.assertTrue(result["success"], result.get("error"))
        self.assertTrue(out.endswith(".ogg"))
        self.assertAlmostEqual(probe_duration(out), 1.0, delta=0.4)

    def test_end_before_start(self):
        result = audio_trim.trim_audio(self.mp3, self.out("bad.mp3"), start=3, end=1)
        self.assertFalse(result["success"])
        self.assertIn("结束时间必须晚于开始时间", result["error"])

    def test_invalid_time_format(self):
        result = audio_trim.trim_audio(self.mp3, self.out("bad2.mp3"), start="十秒")
        self.assertFalse(result["success"])
        self.assertIn("格式不正确", result["error"])

    def test_start_beyond_duration(self):
        result = audio_trim.trim_audio(self.mp3, self.out("far.mp3"), start=99, end=None)
        self.assertFalse(result["success"])
        self.assertIn("这个音频只有", result["error"])

    def test_extra_containers_are_not_rejected(self):
        """扩展名不再做硬性白名单：交给 ffmpeg 判断（aac / opus / wma / aiff 都要能裁）。"""
        for fmt, filename in (("aac", "tone.aac"), ("opus", "tone.opus"),
                              ("wma", "tone.wma"), ("aiff", "tone.aiff"),
                              ("m4a", "tone2.m4a")):
            with self.subTest(fmt=fmt):
                source = make_audio(self.work / filename, duration=3, fmt=fmt)
                out = self.out(f"trimmed-{fmt}.{fmt}")
                result = audio_trim.trim_audio(source, out, start=1, end=2, mode="fast")
                self.assertTrue(result["success"], f"{fmt}: {result.get('error')}")
                self.assertTrue(out.endswith(f".{fmt}"), "必须保留原扩展名")
                self.assertAlmostEqual(probe_duration(out), 1.0, delta=0.5)

    def test_video_file_with_audio_can_be_trimmed(self):
        """以前 .mp4 会被直接拒；现在交给 ffmpeg 判断，视频里的音轨也能裁。"""
        video = make_video(self.work / "clip.mp4", duration=3)
        out = self.out("fromvideo.mp4")
        result = audio_trim.trim_audio(video, out, start=1, end=2, mode="fast")
        self.assertTrue(result["success"], result.get("error"))
        self.assertTrue(out.endswith(".mp4"))
        self.assertIn("audio", stream_types(out))

    def test_unrecognizable_file_reports_common_formats(self):
        garbage = self.work / "broken.xyz"
        garbage.write_bytes(b"this is definitely not audio" * 40)
        result = audio_trim.trim_audio(str(garbage), self.out("broken.mp3"))
        self.assertFalse(result["success"])
        self.assertIn("识别不出", result["error"])
        self.assertIn("mp3", result["error"], "错误里要提示常见格式")
        self.assertIn("音频格式转换", result["error"], "要给出可操作的下一步")

    def test_missing_file(self):
        result = audio_trim.trim_audio(str(self.work / "nope.mp3"), self.out("x.mp3"))
        self.assertFalse(result["success"])
        self.assertIn("找不到", result["error"])

    def test_bad_mode(self):
        result = audio_trim.trim_audio(self.mp3, self.out("m.mp3"), start=0, end=1, mode="warp")
        self.assertFalse(result["success"])
        self.assertIn("裁剪模式只支持", result["error"])

    def test_falls_back_to_reencode_when_copy_fails(self):
        out = self.out("fallback.mp3")
        real_run = audio_trim.run_process

        class FakeResult:
            returncode = 1
            stdout = ""
            stderr = "simulated copy failure"

        def fake_run(cmd, timeout=None):
            if "copy" in cmd:
                return FakeResult()
            return real_run(cmd, timeout=timeout)

        with mock.patch.object(audio_trim, "run_process", fake_run):
            result = audio_trim.trim_audio(self.mp3, out, start=1, end=2, mode="fast")

        self.assertTrue(result["success"], result.get("error"))
        self.assertIn("已自动改用重编码", result["message"])
        self.assertAlmostEqual(probe_duration(out), 1.0, delta=0.35)

    def test_falls_back_when_copy_output_duration_is_wrong(self):
        """流复制"成功"但时长明显不对（mp3 空帧那种情况）也要回退。"""
        out = self.out("fallback2.mp3")
        real_run = audio_trim.run_process

        class FakeResult:
            returncode = 0
            stdout = ""
            stderr = ""

        def fake_run(cmd, timeout=None):
            if "copy" in cmd:
                # 产出一个假的、时长离谱的文件冒充流复制结果
                with open(out, "wb") as fh:
                    fh.write(b"\x00" * 4096)
                return FakeResult()
            return real_run(cmd, timeout=timeout)

        with mock.patch.object(audio_trim, "run_process", fake_run):
            result = audio_trim.trim_audio(self.mp3, out, start=1, end=2, mode="fast")

        self.assertTrue(result["success"], result.get("error"))
        self.assertIn("已自动改用重编码", result["message"])
        self.assertAlmostEqual(probe_duration(out), 1.0, delta=0.35)


# ══════════════════════════════════════════════════════════════════════
# 2b) video-trim 的输出容器（container=same|mp4）
# ══════════════════════════════════════════════════════════════════════
class VideoTrimContainerTests(WorkerToolTest):
    def make(self, ext: str, duration: float = 3, name: str = "clip") -> str:
        """生成指定容器的测试视频（编码按容器挑，尽量贴近真实素材）。"""
        codecs = {
            "mp4": ["-c:v", "libx264", "-preset", "ultrafast", "-g", "10", "-pix_fmt", "yuv420p", "-c:a", "aac"],
            "mkv": ["-c:v", "libx264", "-preset", "ultrafast", "-g", "10", "-pix_fmt", "yuv420p", "-c:a", "aac"],
            "mov": ["-c:v", "libx264", "-preset", "ultrafast", "-g", "10", "-pix_fmt", "yuv420p", "-c:a", "aac"],
            "webm": ["-c:v", "libvpx-vp9", "-crf", "45", "-b:v", "0", "-g", "10", "-c:a", "libopus"],
            "avi": ["-c:v", "libx264", "-preset", "ultrafast", "-g", "10", "-pix_fmt", "yuv420p", "-c:a", "libmp3lame"],
        }[ext]
        path = self.work / f"{name}.{ext}"
        args = [
            ffmpeg_path(), "-y", "-hide_banner", "-loglevel", "error",
            "-f", "lavfi", "-i", f"testsrc=duration={duration}:size=160x120:rate=10",
            "-f", "lavfi", "-i", f"sine=frequency=440:duration={duration}",
        ] + codecs + ["-shortest", str(path)]
        result = run_tool_cmd(args, timeout=120)
        self.assertEqual(result.returncode, 0, f"生成 {ext} 素材失败: {result.stderr[-200:]}")
        return str(path)

    def test_resolve_output_suffix_follows_source_by_default(self):
        for ext in ("mp4", "mkv", "mov", "webm", "avi"):
            with self.subTest(ext=ext):
                src = str(self.work / f"a.{ext}")
                self.assertEqual(video_trim.resolve_output_suffix(src, None), f".{ext}")
                self.assertEqual(video_trim.resolve_output_suffix(src, "same"), f".{ext}")
                self.assertEqual(video_trim.resolve_output_suffix(src, "mp4"), ".mp4")
        # 不在实测白名单里的容器 / 没有扩展名 → 退回 mp4
        self.assertEqual(video_trim.resolve_output_suffix(str(self.work / "a.mpg"), "same"), ".mp4")
        self.assertEqual(video_trim.resolve_output_suffix(str(self.work / "a.mpeg"), "same"), ".mp4")
        self.assertEqual(video_trim.resolve_output_suffix(str(self.work / "noext"), "same"), ".mp4")
        self.assertEqual(video_trim.resolve_output_suffix(str(self.work / "a.xyz"), "same"), ".mp4")

    def test_normalize_container_is_lenient(self):
        for value in (None, "", "same", "SAME", "Same", "源", "keep", "garbage", "mkv", 12):
            with self.subTest(value=value):
                self.assertEqual(video_trim.normalize_container(value), "same")
        for value in ("mp4", "MP4", ".mp4", " Mp4 "):
            with self.subTest(value=value):
                self.assertEqual(video_trim.normalize_container(value), "mp4")

    def test_mkv_source_keeps_mkv_container_and_can_stream_copy(self):
        source = self.make("mkv")
        out = str(self.work / "clip-trimmed.mkv")
        result = video_trim.trim_video(source, out, start=1, end=3, mode="fast")
        self.assertTrue(result["success"], result.get("error"))
        self.assertEqual(result["mime"], "video/x-matroska")
        self.assertEqual(result["container"], "mkv")
        self.assertIn("流复制", result["message"])
        self.assertIn(".mkv", result["message"])
        self.assertAlmostEqual(probe_duration(out), 2.0, delta=0.6)

    def test_webm_source_fast_and_precise(self):
        source = self.make("webm", duration=3)
        fast_out = str(self.work / "webm-fast.webm")
        fast = video_trim.trim_video(source, fast_out, start=1, end=2, mode="fast")
        self.assertTrue(fast["success"], fast.get("error"))
        self.assertEqual(fast["mime"], "video/webm")
        self.assertIn("流复制", fast["message"])

        precise_out = str(self.work / "webm-precise.webm")
        precise = video_trim.trim_video(source, precise_out, start=1, end=2, mode="precise")
        self.assertTrue(precise["success"], precise.get("error"))
        self.assertEqual(precise["mime"], "video/webm")
        self.assertIn(".webm", precise["message"])
        self.assertAlmostEqual(probe_duration(precise_out), 1.0, delta=0.5)

    def test_mov_and_avi_sources(self):
        for ext in ("mov", "avi"):
            with self.subTest(ext=ext):
                source = self.make(ext)
                fast_out = str(self.work / f"trim-{ext}.{ext}")
                fast = video_trim.trim_video(source, fast_out, start=1, end=3, mode="fast")
                self.assertTrue(fast["success"], fast.get("error"))
                self.assertEqual(fast["container"], ext)
                # 流复制的切点会对齐关键帧/音频帧，时长只能放宽验证（不能短于请求、不能离谱）
                fast_dur = probe_duration(fast_out)
                self.assertGreaterEqual(fast_dur, 0.9, "不应该比请求的片段还短")
                self.assertLessEqual(fast_dur, 3.4, "不应该把整段都带出来")

                precise_out = str(self.work / f"trim-{ext}-precise.{ext}")
                precise = video_trim.trim_video(source, precise_out, start=1, end=3, mode="precise")
                self.assertTrue(precise["success"], precise.get("error"))
                self.assertEqual(precise["container"], ext)
                self.assertAlmostEqual(probe_duration(precise_out), 2.0, delta=0.4)

    def test_container_mp4_forces_mp4_and_says_so(self):
        source = self.make("mkv")
        out = str(self.work / "forced.mp4")
        result = video_trim.trim_video(source, out, start=1, end=3, mode="precise",
                                       container="mp4")
        self.assertTrue(result["success"], result.get("error"))
        self.assertEqual(result["mime"], "video/mp4")
        self.assertIn("强制输出为 mp4", result["message"])
        self.assertAlmostEqual(probe_duration(out), 2.0, delta=0.4)

    def test_unsupported_source_extension_falls_back_to_mp4_with_note(self):
        source = self.make("mkv", name="weird")
        renamed = self.work / "weird.xyz"
        source_path = Path(source)
        source_path.replace(renamed)  # 换成 ffmpeg 认不出容器的扩展名（内容仍是 mkv）
        suffix = video_trim.resolve_output_suffix(str(renamed), "same")
        self.assertEqual(suffix, ".mp4")
        out = str(self.work / f"weird-trimmed{suffix}")
        result = video_trim.trim_video(str(renamed), out, start=1, end=2, mode="fast")
        self.assertTrue(result["success"], result.get("error"))
        self.assertIn("不能直接当输出容器", result["message"])

    def test_precise_falls_back_to_mp4_when_container_refuses_reencode(self):
        """源容器不接受重编码时，退回 mp4 再试，而不是直接把任务判失败。"""
        source = self.make("mkv")
        out = str(self.work / "refuse.mkv")
        real_run = video_trim.run_process

        class FakeResult:
            returncode = 1
            stdout = ""
            stderr = "simulated: container refuses reencode"

        def fake_run(cmd, timeout=None):
            if str(cmd[-1]).endswith(".mkv"):
                return FakeResult()
            return real_run(cmd, timeout=timeout)

        with mock.patch.object(video_trim, "run_process", fake_run):
            result = video_trim.trim_video(source, out, start=1, end=2, mode="precise")

        self.assertTrue(result["success"], result.get("error"))
        self.assertTrue(result["output"].endswith(".mp4"))
        self.assertIn("已改用 mp4 输出", result["message"])
        self.assertEqual(result["mime"], "video/mp4")

    def test_job_keeps_container_and_mime(self):
        source = self.make("mkv")
        job = self.run_job("video-trim", {"file": source, "start": "1", "end": "2", "mode": "fast"})
        path = self.assert_job_ok(job)
        self.assertEqual(job["resultFilename"], "clip-trimmed.mkv")
        self.assertEqual(job["resultMimeType"], "video/x-matroska")
        self.assertTrue(path.is_file())

    def test_job_with_chinese_file_name_keeps_extension(self):
        """中文文件名 + 中文交付名：扩展名必须活下来（_safe_filename 修复后的端到端证据）。"""
        source = self.make("mp4", name="我的视频")
        job = self.run_job("video-trim", {"file": source, "start": "1", "end": "2", "mode": "precise"})
        path = self.assert_job_ok(job)
        self.assertEqual(job["resultFilename"], "我的视频-trimmed.mp4")
        self.assertTrue(path.name.endswith("我的视频-trimmed.mp4"))


# ══════════════════════════════════════════════════════════════════════
# 3) video-frame-extract（本地抽帧；注意与既有 video-thumbnail 在线封面下载不同）
# ══════════════════════════════════════════════════════════════════════
class VideoFrameExtractTests(WorkerToolTest):
    def setUp(self):
        super().setUp()
        self.video = make_video(self.work / "clip.mp4", duration=5, gop=10)
        self.audio = make_audio(self.work / "tone.mp3", duration=2)

    def test_png_frame_default_time(self):
        out = str(self.work / "frame.png")
        result = video_trim.extract_thumbnail(self.video, out, time=None, fmt="png")
        self.assertTrue(result["success"], result.get("error"))
        with Image.open(out) as img:
            self.assertEqual(img.format, "PNG")
            self.assertEqual(img.size, (320, 240))
        self.assertEqual(result["mime"], "image/png")

    def test_jpg_frame_with_width_scaling(self):
        out = str(self.work / "frame.jpg")
        result = video_trim.extract_thumbnail(self.video, out, time="1.5", fmt="jpg", width="160")
        self.assertTrue(result["success"], result.get("error"))
        with Image.open(out) as img:
            self.assertEqual(img.format, "JPEG")
            self.assertEqual(img.size, (160, 120))
        self.assertEqual(result["mime"], "image/jpeg")

    def test_requested_time_actually_changes_the_frame(self):
        first = str(self.work / "t0.png")
        later = str(self.work / "t15.png")
        self.assertTrue(video_trim.extract_thumbnail(self.video, first, time="0", fmt="png")["success"])
        self.assertTrue(video_trim.extract_thumbnail(self.video, later, time="1.5", fmt="png")["success"])
        self.assertNotEqual(
            Path(first).read_bytes(), Path(later).read_bytes(),
            "两个时间点抽出来的画面居然一模一样，说明 time 参数没生效",
        )

    def test_hhmmss_time_accepted(self):
        out = str(self.work / "hhmmss.png")
        result = video_trim.extract_thumbnail(self.video, out, time="00:00:02.5", fmt="png")
        self.assertTrue(result["success"], result.get("error"))
        self.assertIn("2.5", result["message"])

    def test_time_beyond_duration_reports_total_length(self):
        out = str(self.work / "far.png")
        result = video_trim.extract_thumbnail(self.video, out, time="30", fmt="png")
        self.assertFalse(result["success"])
        self.assertIn("这个视频只有", result["error"])
        self.assertIn("取不到第 30 秒的画面", result["error"])
        self.assertFalse(os.path.exists(out))

    def test_bad_format_is_rejected(self):
        result = video_trim.extract_thumbnail(self.video, str(self.work / "f.webp"), fmt="webp")
        self.assertFalse(result["success"])
        self.assertIn("仅支持 png 或 jpg", result["error"])

    def test_bad_time_format(self):
        result = video_trim.extract_thumbnail(self.video, str(self.work / "f.png"), time="一秒")
        self.assertFalse(result["success"])
        self.assertIn("格式不正确", result["error"])

    def test_bad_width(self):
        for bad, needle in (("abc", "必须是数字"), ("0", "至少为 2"), ("99999", "太大")):
            with self.subTest(bad=bad):
                result = video_trim.extract_thumbnail(self.video, str(self.work / "w.png"), width=bad)
                self.assertFalse(result["success"])
                self.assertIn(needle, result["error"])

    def test_audio_file_is_rejected(self):
        result = video_trim.extract_thumbnail(self.audio, str(self.work / "a.png"))
        self.assertFalse(result["success"])
        self.assertIn("没有视频画面", result["error"])

    def test_missing_file(self):
        result = video_trim.extract_thumbnail(str(self.work / "nope.mp4"), str(self.work / "n.png"))
        self.assertFalse(result["success"])
        self.assertIn("找不到", result["error"])


# ══════════════════════════════════════════════════════════════════════
# 4) csv-excel
# ══════════════════════════════════════════════════════════════════════
CHINESE_CSV = (
    "姓名,备注,编号,身份证号,数量,单价,公式\n"
    "张三,\"含,逗号\",007,110101199003074512,12,3.5,=1+1\n"
    "李四,\"含\"\"引号\"\"\",00123,123456789012345678,0,10.25,=SUM(A1:A2)\n"
    "王五,\"第一行\n第二行\",0999,1234567890123456,-5,0.5,普通文本\n"
)


def write_csv(path, text: str, encoding: str = "utf-8", newline: str = "") -> str:
    path.write_text(text, encoding=encoding, newline=newline)
    return str(path)


class CsvToXlsxTests(WorkerToolTest):
    def convert(self, csv_path, direction="to-xlsx", **kwargs):
        return csv_excel.convert_file(csv_path, self.work / "out", direction=direction, **kwargs)

    def test_chinese_special_chars_and_cell_types(self):
        csv = write_csv(self.work / "表格.csv", CHINESE_CSV)
        result = self.convert(csv)
        self.assertTrue(result["success"], result.get("error"))
        self.assertTrue(result["output"].endswith(".xlsx"))

        wb = load_workbook(result["output"])
        ws = wb.active

        # 表头加粗 + 冻结首行 + 列宽在合理区间
        self.assertEqual(ws["A1"].value, "姓名")
        self.assertTrue(ws["A1"].font.bold, "表头应该加粗")
        self.assertEqual(ws.freeze_panes, "A2", "首行应该被冻结")
        for column in ("A", "B", "C", "D", "E", "F", "G"):
            width = ws.column_dimensions[column].width
            self.assertGreaterEqual(width, 8.0)
            self.assertLessEqual(width, 60.0)

        # 中文 / 逗号 / 引号 / 换行都原样保留在一个单元格里
        self.assertEqual(ws["A2"].value, "张三")
        self.assertEqual(ws["B2"].value, "含,逗号")
        self.assertEqual(ws["B3"].value, '含"引号"')
        self.assertEqual(str(ws["B4"].value).replace("\r\n", "\n"), "第一行\n第二行")

        # 前导 0 的编号 / 超长数字必须是文本（否则 Excel 会破坏它们）
        self.assertEqual(ws["C2"].value, "007")
        self.assertEqual(ws["C2"].data_type, "s")
        self.assertEqual(ws["C3"].value, "00123")
        self.assertEqual(ws["C3"].data_type, "s")
        self.assertEqual(ws["C4"].value, "0999")
        self.assertEqual(ws["D2"].value, "110101199003074512")
        self.assertEqual(ws["D2"].data_type, "s", "18 位身份证号必须保持文本")
        self.assertEqual(ws["D3"].value, "123456789012345678")
        self.assertEqual(ws["D3"].data_type, "s")
        self.assertEqual(ws["D4"].value, "1234567890123456")
        self.assertEqual(ws["D4"].data_type, "s", "16 位长数字必须保持文本")

        # 普通数字必须是真数字，能在 Excel 里直接求和
        self.assertEqual(ws["E2"].value, 12)
        self.assertEqual(ws["E2"].data_type, "n")
        self.assertEqual(ws["E4"].value, -5)
        self.assertEqual(ws["F2"].value, 3.5)
        self.assertEqual(ws["F4"].value, 0.5)
        self.assertEqual(ws["F4"].data_type, "n")

        # = 开头的内容不能被当成公式（CSV 注入）
        self.assertEqual(ws["G2"].value, "=1+1")
        self.assertEqual(ws["G2"].data_type, "s", "= 开头必须按文本写，不能变成公式")
        self.assertEqual(ws["G3"].data_type, "s")
        self.assertEqual(ws["G4"].value, "普通文本")
        wb.close()

    def test_gbk_encoded_csv(self):
        csv = write_csv(self.work / "gbk.csv", "姓名,城市\n张三,北京\n", encoding="gb18030")
        result = self.convert(csv)
        self.assertTrue(result["success"], result.get("error"))
        self.assertIn("gb18030", result["message"])
        wb = load_workbook(result["output"])
        ws = wb.active
        self.assertEqual(ws["A2"].value, "张三")
        self.assertEqual(ws["B2"].value, "北京")
        wb.close()

    def test_utf8_bom_is_stripped(self):
        csv = write_csv(self.work / "bom.csv", "姓名,城市\n张三,北京\n", encoding="utf-8-sig")
        result = self.convert(csv)
        self.assertTrue(result["success"], result.get("error"))
        wb = load_workbook(result["output"])
        ws = wb.active
        self.assertEqual(ws["A1"].value, "姓名", "BOM 必须被去掉，不能粘在表头里")
        wb.close()

    def test_semicolon_delimiter_detected(self):
        csv = write_csv(self.work / "semi.csv", "姓名;年龄;城市\n张三;30;北京\n李四;25;上海\n")
        result = self.convert(csv)
        self.assertTrue(result["success"], result.get("error"))
        self.assertIn("分号", result["message"])
        wb = load_workbook(result["output"])
        ws = wb.active
        self.assertEqual(ws["A1"].value, "姓名")
        self.assertEqual(ws["B1"].value, "年龄")
        self.assertEqual(ws["C1"].value, "城市")
        self.assertEqual(ws["B2"].value, 30)
        wb.close()

    def test_tab_delimiter_detected(self):
        csv = write_csv(self.work / "tab.csv", "a\tb\tc\n1\t2\t3\n4\t5\t6\n")
        result = self.convert(csv)
        self.assertTrue(result["success"], result.get("error"))
        self.assertIn("Tab", result["message"])
        wb = load_workbook(result["output"])
        ws = wb.active
        self.assertEqual([ws.cell(row=1, column=i).value for i in (1, 2, 3)], ["a", "b", "c"])
        wb.close()

    def test_pipe_delimiter_detected(self):
        csv = write_csv(self.work / "pipe.csv", "a|b|c\n1|2|3\n4|5|6\n")
        result = self.convert(csv)
        self.assertTrue(result["success"], result.get("error"))
        self.assertIn("竖线", result["message"])
        wb = load_workbook(result["output"])
        self.assertEqual(wb.active.max_column, 3)
        wb.close()

    def test_explicit_delimiter_overrides_detection(self):
        csv = write_csv(self.work / "explicit.csv", "a;b;c\n1;2;3\n")
        result = self.convert(csv, delimiter=";")
        self.assertTrue(result["success"], result.get("error"))
        wb = load_workbook(result["output"])
        self.assertEqual(wb.active.max_column, 3)
        wb.close()

        result2 = self.convert(csv, delimiter=",")
        self.assertTrue(result2["success"], result2.get("error"))
        wb2 = load_workbook(result2["output"])
        self.assertEqual(wb2.active.max_column, 1, "显式指定逗号时应该只有一列")
        wb2.close()

    def test_bad_delimiter_reports_human_error(self):
        csv = write_csv(self.work / "bad.csv", "a,b\n1,2\n")
        result = self.convert(csv, delimiter=";;;")
        self.assertFalse(result["success"])
        self.assertIn("分隔符只能是单个字符", result["error"])

    def test_empty_lines_and_ragged_rows(self):
        csv = write_csv(self.work / "ragged.csv", "a,b,c\n1,2\n\n4,5,6\n")
        result = self.convert(csv)
        self.assertTrue(result["success"], result.get("error"))
        wb = load_workbook(result["output"])
        ws = wb.active
        self.assertEqual(ws["A2"].value, 1)
        self.assertIsNone(ws["C2"].value, "缺列的地方应该是空的")
        self.assertIsNone(ws["A3"].value, "空行要保持空行")
        self.assertEqual(ws["A4"].value, 4, "空行之后的第四行不能被挤到第三行")
        self.assertEqual(ws["C4"].value, 6)
        wb.close()

    def test_header_shorter_than_data_rows(self):
        csv = write_csv(self.work / "short-header.csv", "a\n1,2,3\n")
        result = self.convert(csv)
        self.assertTrue(result["success"], result.get("error"))
        wb = load_workbook(result["output"])
        ws = wb.active
        self.assertEqual(ws["A1"].value, "a")
        self.assertEqual(ws["B1"].value, "列2")
        self.assertEqual(ws["C1"].value, "列3")
        wb.close()

    def test_sheet_name_from_payload(self):
        csv = write_csv(self.work / "s.csv", "a,b\n1,2\n")
        result = self.convert(csv, sheet="数据表")
        self.assertTrue(result["success"], result.get("error"))
        wb = load_workbook(result["output"])
        self.assertEqual(wb.active.title, "数据表")
        wb.close()

    def test_illegal_sheet_name_is_sanitized(self):
        csv = write_csv(self.work / "s2.csv", "a,b\n1,2\n")
        result = self.convert(csv, sheet="a:b*c?")
        self.assertTrue(result["success"], result.get("error"))
        wb = load_workbook(result["output"])
        self.assertNotIn(":", wb.active.title)
        wb.close()

    def test_very_long_text_is_truncated_in_header_but_kept_in_cells(self):
        long_header = "很长的表头" * 60
        csv = write_csv(self.work / "long.csv", f"{long_header}\n值\n")
        result = self.convert(csv)
        self.assertTrue(result["success"], result.get("error"))
        wb = load_workbook(result["output"])
        ws = wb.active
        self.assertLessEqual(len(str(ws["A1"].value)), 201)
        self.assertLessEqual(ws.column_dimensions["A"].width, 60.0)
        wb.close()

    def test_row_limit_reports_human_error(self):
        csv = write_csv(self.work / "many.csv", "a,b\n" + "1,2\n" * 10)
        original = csv_excel.MAX_CSV_ROWS
        csv_excel.MAX_CSV_ROWS = 3
        try:
            result = self.convert(csv)
        finally:
            csv_excel.MAX_CSV_ROWS = original
        self.assertFalse(result["success"])
        self.assertIn("行数太多", result["error"])

    def test_empty_csv_is_rejected(self):
        csv = write_csv(self.work / "empty.csv", "   \n\n")
        result = self.convert(csv)
        self.assertFalse(result["success"])
        self.assertIn("CSV 文件是空的", result["error"])

    def test_unknown_extension_with_auto_direction(self):
        txt = write_csv(self.work / "note.txt", "a,b\n1,2\n")
        result = self.convert(txt, direction="auto")
        self.assertFalse(result["success"])
        self.assertIn("无法根据扩展名判断转换方向", result["error"])

    def test_explicit_direction_allows_txt(self):
        txt = write_csv(self.work / "note2.txt", "a,b\n1,2\n")
        result = self.convert(txt, direction="to-xlsx")
        self.assertTrue(result["success"], result.get("error"))
        wb = load_workbook(result["output"])
        self.assertEqual(wb.active.max_column, 2)
        wb.close()

    def test_bad_direction(self):
        csv = write_csv(self.work / "d.csv", "a,b\n1,2\n")
        result = self.convert(csv, direction="sideways")
        self.assertFalse(result["success"])
        self.assertIn("无法识别的转换方向", result["error"])

    def test_missing_file(self):
        result = self.convert(str(self.work / "nope.csv"))
        self.assertFalse(result["success"])
        self.assertIn("找不到文件", result["error"])


class CsvHasHeaderTests(WorkerToolTest):
    """payload 里的 has_header（首行是否为表头）：true=加粗+冻结首行，false=当普通数据。"""

    CSV = "姓名,年龄\n张三,30\n李四,25\n"

    def convert(self, csv_path, **kwargs):
        return csv_excel.convert_file(csv_path, self.work / "out", direction="to-xlsx", **kwargs)

    def test_has_header_true_bolds_and_freezes_first_row(self):
        csv = write_csv(self.work / "h_true.csv", self.CSV)
        result = self.convert(csv, has_header="true")
        self.assertTrue(result["success"], result.get("error"))
        self.assertIn("首行作表头", result["message"])
        wb = load_workbook(result["output"])
        ws = wb.active
        self.assertTrue(ws["A1"].font.bold, "has_header=true 时首行必须加粗")
        self.assertTrue(ws["B1"].font.bold)
        self.assertEqual(ws.freeze_panes, "A2", "has_header=true 时必须冻结首行")
        self.assertEqual(ws["A1"].value, "姓名")
        self.assertEqual(ws["A2"].value, "张三")
        wb.close()

    def test_has_header_false_keeps_first_row_as_plain_data(self):
        csv = write_csv(self.work / "h_false.csv", self.CSV)
        result = self.convert(csv, has_header="false")
        self.assertTrue(result["success"], result.get("error"))
        self.assertIn("首行按普通数据", result["message"])
        wb = load_workbook(result["output"])
        ws = wb.active
        self.assertFalse(ws["A1"].font.bold, "has_header=false 时首行不能加粗")
        self.assertFalse(ws["B1"].font.bold)
        self.assertIsNone(ws.freeze_panes, "has_header=false 时不该冻结首行")
        self.assertEqual(ws["A1"].value, "姓名", "首行要原样当数据保留")
        self.assertEqual(ws["A2"].value, "张三")
        self.assertEqual(ws["A3"].value, "李四")
        # 列宽照常按内容估算
        self.assertGreaterEqual(ws.column_dimensions["A"].width, 8.0)
        wb.close()

    def test_default_is_header(self):
        csv = write_csv(self.work / "h_default.csv", self.CSV)
        result = self.convert(csv)
        self.assertTrue(result["success"], result.get("error"))
        self.assertIn("首行作表头", result["message"])
        wb = load_workbook(result["output"])
        self.assertTrue(wb.active["A1"].font.bold)
        wb.close()

    def test_truthy_and_falsy_spellings(self):
        csv = write_csv(self.work / "h_spell.csv", self.CSV)
        for value, expect_bold in (
            ("true", True), ("TRUE", True), ("1", True), ("yes", True), ("Yes", True),
            (" True ", True), (True, True), (1, True),
            ("false", False), ("False", False), ("0", False), ("no", False), ("NO", False),
            (False, False), (0, False),
        ):
            with self.subTest(value=value):
                safe = re.sub(r"[^0-9a-zA-Z]+", "_", str(value)) or "none"
                out_dir = self.work / f"out_{safe}"
                result = csv_excel.convert_file(csv, out_dir, direction="to-xlsx", has_header=value)
                self.assertTrue(result["success"], result.get("error"))
                wb = load_workbook(result["output"])
                self.assertEqual(wb.active["A1"].font.bold, expect_bold,
                                 f"has_header={value!r} 的加粗行为不对")
                wb.close()

    def test_invalid_value_falls_back_to_true_without_error(self):
        csv = write_csv(self.work / "h_bad.csv", self.CSV)
        for value in ("maybe", "表头是啥", "", None, [], {}):
            with self.subTest(value=value):
                result = csv_excel.convert_file(csv, self.work / "out_bad", direction="to-xlsx",
                                               has_header=value)
                self.assertTrue(result["success"], f"has_header={value!r} 不该让转换失败: {result.get('error')}")
                wb = load_workbook(result["output"])
                self.assertTrue(wb.active["A1"].font.bold)
                wb.close()

    def test_has_header_false_still_writes_real_numbers(self):
        csv = write_csv(self.work / "h_num.csv", "编号,数量\n007,12\n00123,30\n")
        result = self.convert(csv, has_header="false")
        self.assertTrue(result["success"], result.get("error"))
        wb = load_workbook(result["output"])
        ws = wb.active
        self.assertEqual(ws["A1"].value, "编号")
        self.assertEqual(ws["A2"].value, "007")
        self.assertEqual(ws["A2"].data_type, "s", "编号仍要保持文本")
        self.assertEqual(ws["B2"].value, 12)
        self.assertEqual(ws["B2"].data_type, "n")
        wb.close()

    def test_xlsx_to_csv_accepts_has_header_without_surprises(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["姓名", "年龄"])
        ws.append(["张三", 30])
        xlsx = self.work / "h.xlsx"
        wb.save(str(xlsx))
        wb.close()

        with_header = csv_excel.convert_file(xlsx, self.work / "o1", direction="to-csv", has_header="true")
        without = csv_excel.convert_file(xlsx, self.work / "o2", direction="to-csv", has_header="false")
        self.assertTrue(with_header["success"], with_header.get("error"))
        self.assertTrue(without["success"], without.get("error"))
        self.assertEqual(Path(with_header["output"]).read_bytes(), Path(without["output"]).read_bytes(),
                         "导出方向不对首行做特殊处理，两种取值的产物应当一致")
        self.assertIn("姓名,年龄", Path(with_header["output"]).read_text(encoding="utf-8-sig"))


class XlsxToCsvTests(WorkerToolTest):
    def make_workbook(self) -> str:
        wb = Workbook()
        first = wb.active
        first.title = "Sheet1"
        first.append(["姓名", "年龄", "备注"])
        first.append(["张三", 30, "含,逗号"])
        first.append(["李四", 25.5, '含"引号"'])
        second = wb.create_sheet("第二张表")
        second.append(["列A", "列B"])
        second.append(["多行\n内容", "=SUM(A1)"])
        path = self.work / "book.xlsx"
        wb.save(str(path))
        wb.close()
        return str(path)

    def convert(self, xlsx, direction="to-csv", **kwargs):
        return csv_excel.convert_file(xlsx, self.work / "out", direction=direction, **kwargs)

    def test_first_sheet_by_default(self):
        result = self.convert(self.make_workbook())
        self.assertTrue(result["success"], result.get("error"))
        raw = Path(result["output"]).read_bytes()
        self.assertTrue(raw.startswith(b"\xef\xbb\xbf"), "CSV 必须是 UTF-8 带 BOM（Excel 打开不乱码）")
        text = raw.decode("utf-8-sig")
        self.assertIn("姓名,年龄,备注", text)
        self.assertIn("张三,30,", text)

    def test_sheet_selection(self):
        result = self.convert(self.make_workbook(), sheet="第二张表")
        self.assertTrue(result["success"], result.get("error"))
        self.assertIn("第二张表", result["message"])
        text = Path(result["output"]).read_bytes().decode("utf-8-sig")
        self.assertIn("列A", text)
        self.assertNotIn("姓名", text)

    def test_missing_sheet_lists_available_names(self):
        result = self.convert(self.make_workbook(), sheet="不存在的表")
        self.assertFalse(result["success"])
        self.assertIn("找不到工作表", result["error"])
        self.assertIn("Sheet1", result["error"])
        self.assertIn("第二张表", result["error"])

    def test_special_chars_roundtrip(self):
        import csv as csv_module
        import io as io_module

        wb = Workbook()
        ws = wb.active
        ws.append(["含,逗号", '含"引号"', "含\n换行", "普通"])
        path = self.work / "special.xlsx"
        wb.save(str(path))
        wb.close()

        result = self.convert(str(path))
        self.assertTrue(result["success"], result.get("error"))
        text = Path(result["output"]).read_bytes().decode("utf-8-sig")
        rows = list(csv_module.reader(io_module.StringIO(text)))
        self.assertEqual(rows[0][0], "含,逗号")
        self.assertEqual(rows[0][1], '含"引号"')
        self.assertEqual(rows[0][2], "含\n换行")
        self.assertEqual(rows[0][3], "普通")
        # 真 CSV 里逗号字段必须被引号包起来
        self.assertIn('"含,逗号"', text)

    def test_number_and_date_formatting(self):
        wb = Workbook()
        ws = wb.active
        ws.append([10.0, 3.5, True, "00123"])
        path = self.work / "types.xlsx"
        wb.save(str(path))
        wb.close()
        result = self.convert(str(path))
        self.assertTrue(result["success"], result.get("error"))
        text = Path(result["output"]).read_bytes().decode("utf-8-sig").strip()
        self.assertEqual(text, "10,3.5,TRUE,00123")

    def test_not_an_excel_file_reports_human_error(self):
        fake = write_csv(self.work / "fake.xlsx", "a,b\n1,2\n")
        result = self.convert(fake)
        self.assertFalse(result["success"])
        self.assertIn("无法读取 Excel 文件", result["error"])

    def test_empty_sheet_is_rejected(self):
        wb = Workbook()
        path = self.work / "empty.xlsx"
        wb.save(str(path))
        wb.close()
        result = self.convert(str(path))
        self.assertFalse(result["success"])
        self.assertIn("没有任何内容", result["error"])

    def test_missing_file(self):
        result = self.convert(str(self.work / "nope.xlsx"))
        self.assertFalse(result["success"])
        self.assertIn("找不到文件", result["error"])

    def test_direction_auto_on_xlsx(self):
        result = csv_excel.convert_file(self.make_workbook(), self.work / "out", direction="auto")
        self.assertTrue(result["success"], result.get("error"))
        self.assertTrue(result["output"].endswith(".csv"))


# ══════════════════════════════════════════════════════════════════════
# 5) markdown-to-pdf
# ══════════════════════════════════════════════════════════════════════
MD_SAMPLE = """# 中文一级标题 FurinaKit

这是一段中文正文，里面还有 English words 和数字 12345，用来验证中英混排断行不会溢出页面宽度。

## 二级标题

**加粗中文** 与 *斜体中文* 与 `行内代码` 混排。

### 三级标题

- 无序列表第一项
- 无序列表第二项

  1. 嵌套有序第一项
  2. 嵌套有序第二项

> 引用第一行
> 引用第二行

```python
def 中文函数(参数):
    return "hello 世界"
```

| 姓名 | 年龄 | 备注 |
| --- | ---: | :---: |
| 张三 | 30 | 中文很长的单元格内容测试 |
| Bob | 5 | x |

---

1. 有序第一项
2. 有序第二项

<script>alert(1)</script>
"""

MD_PROBES = [
    "中文一级标题",
    "这是一段中文正文",
    "二级标题",
    "加粗中文",
    "无序列表第一项",
    "嵌套有序第一项",
    "嵌套有序第二项",
    "引用第一行",
    "引用第二行",
    "def 中文函数",
    "hello 世界",
    "中文很长的单元格内容测试",
    "有序第一项",
    "有序第二项",
]


class MarkdownToPdfTests(WorkerToolTest):
    def render(self, md: str, name: str = "document.pdf", **kwargs):
        out = self.work / name
        result = markdown_pdf.markdown_to_pdf(text=md, output_path=str(out), **kwargs)
        self.assertTrue(result["success"], result.get("error"))
        self.assertTrue(out.is_file())
        return out, result

    def test_chinese_text_is_real_extractable_text(self):
        out, result = self.render(MD_SAMPLE)
        self.assertEqual(result["mime"], "application/pdf")
        self.assertGreaterEqual(result["pages"], 1)
        raw = out.read_bytes()
        self.assertTrue(raw.startswith(b"%PDF"), "输出必须是真 PDF")

        doc = fitz.open(str(out))
        try:
            text = "\n".join(page.get_text() for page in doc)
        finally:
            doc.close()
        for probe in MD_PROBES:
            self.assertIn(probe, text, f"PDF 文本里读不到「{probe}」，中文可能没被正确写入")

    def test_code_block_chinese_is_readable(self):
        out, _ = self.render(MD_SAMPLE)
        doc = fitz.open(str(out))
        try:
            text = "\n".join(page.get_text() for page in doc)
        finally:
            doc.close()
        self.assertIn("def 中文函数", text)
        self.assertIn("hello 世界", text)

    def test_html_is_rendered_as_plain_text(self):
        out, _ = self.render(MD_SAMPLE)
        doc = fitz.open(str(out))
        try:
            text = "\n".join(page.get_text() for page in doc)
        finally:
            doc.close()
        self.assertIn("alert(1)", text, "HTML 应该被当纯文本画出来")
        links = fitz.open(str(out))
        try:
            for page in links:
                self.assertEqual(len(page.get_links()), 0, "不该凭空出现可点击链接")
        finally:
            links.close()

    def test_heading_is_larger_than_body_text(self):
        out, _ = self.render(MD_SAMPLE)
        doc = fitz.open(str(out))
        try:
            sizes = {}
            for page in doc:
                for block in page.get_text("dict")["blocks"]:
                    for line in block.get("lines", []):
                        for span in line["spans"]:
                            sizes[span["text"]] = span["size"]
        finally:
            doc.close()
        heading = next(size for text, size in sizes.items() if text.startswith("中文一级标题"))
        body = next(size for text, size in sizes.items() if text.startswith("这是一段中文正文"))
        third = next(size for text, size in sizes.items() if text.startswith("三级标题"))
        second = next(size for text, size in sizes.items() if text.startswith("二级标题"))
        self.assertGreater(heading, second)
        self.assertGreater(second, third)
        self.assertGreater(third, body)

    def test_tables_and_rules_are_drawn(self):
        out, _ = self.render(MD_SAMPLE)
        doc = fitz.open(str(out))
        try:
            drawings = doc[0].get_drawings()
        finally:
            doc.close()
        self.assertGreaterEqual(len(drawings), 4, "表格线 / 分割线 / 引用竖线应该被画出来")

    def test_nothing_overflows_the_page(self):
        for md, label in ((MD_SAMPLE, "样例"), ("x" * 400, "超长英文"), ("中" * 2000, "超长中文"),
                          ("word " * 500, "超长句子")):
            with self.subTest(case=label):
                out, _ = self.render(md, name=f"overflow_{label}.pdf")
                doc = fitz.open(str(out))
                try:
                    for index, page in enumerate(doc, 1):
                        rect = page.rect
                        for block in page.get_text("blocks"):
                            x0, y0, x1, y1 = block[:4]
                            self.assertGreaterEqual(x0, -1, f"{label} 第 {index} 页有内容跑到左边外面")
                            self.assertLessEqual(x1, rect.width + 1, f"{label} 第 {index} 页有内容溢出右边")
                            self.assertGreaterEqual(y0, -1, f"{label} 第 {index} 页有内容跑到上边外面")
                            self.assertLessEqual(y1, rect.height + 1, f"{label} 第 {index} 页有内容溢出下边")
                finally:
                    doc.close()

    def test_footer_page_numbers(self):
        out, result = self.render(MD_SAMPLE)
        doc = fitz.open(str(out))
        try:
            text = doc[0].get_text()
            total = doc.page_count
        finally:
            doc.close()
        self.assertIn(f"第 1 页 / 共 {total} 页", text)

    def test_multipage_document_has_footer_on_every_page(self):
        md = "\n\n".join(f"## 第{i}段中文标题\n\n这是第{i}段中文正文，用来把文档撑到多页。" for i in range(1, 120))
        out, result = self.render(md, name="long.pdf")
        self.assertGreaterEqual(result["pages"], 2)
        doc = fitz.open(str(out))
        try:
            total = doc.page_count
            for index in range(total):
                self.assertIn(f"第 {index + 1} 页 / 共 {total} 页", doc[index].get_text())
        finally:
            doc.close()

    def test_long_table_breaks_across_pages_without_overflow(self):
        rows = "\n".join(f"| 行{i} | {i} | 中文内容{i}自动换行测试 |" for i in range(1, 90))
        md = "| 姓名 | 序号 | 备注 |\n| --- | --- | --- |\n" + rows + "\n"
        out, result = self.render(md, name="longtable.pdf")
        self.assertGreaterEqual(result["pages"], 2)
        doc = fitz.open(str(out))
        try:
            for index, page in enumerate(doc, 1):
                rect = page.rect
                for block in page.get_text("blocks"):
                    self.assertLessEqual(block[2], rect.width + 1, f"第 {index} 页表格溢出右边")
                    self.assertLessEqual(block[3], rect.height + 1, f"第 {index} 页表格溢出下边")
            pages_with_header = sum(1 for page in doc if "姓名" in page.get_text())
            self.assertGreaterEqual(pages_with_header, 2, "换页后应该重画表头")
        finally:
            doc.close()

    def test_letter_page_size(self):
        out, _ = self.render(MD_SAMPLE, name="letter.pdf", page_size="letter")
        doc = fitz.open(str(out))
        try:
            width, height = doc[0].rect.width, doc[0].rect.height
        finally:
            doc.close()
        self.assertAlmostEqual(width, 612.0, delta=1.0)
        self.assertAlmostEqual(height, 792.0, delta=1.0)

    def test_font_size_is_applied(self):
        out, _ = self.render(MD_SAMPLE, name="fs20.pdf", font_size="20")
        doc = fitz.open(str(out))
        try:
            sizes = [span["size"] for block in doc[0].get_text("dict")["blocks"]
                     for line in block.get("lines", []) for span in line["spans"]]
        finally:
            doc.close()
        self.assertTrue(any(size >= 19.5 for size in sizes), f"字号没有生效: {sorted(set(sizes))}")

    def test_bad_font_size(self):
        for bad, needle in ((100, "字号必须在"), ("abc", "必须是数字")):
            with self.subTest(bad=bad):
                result = markdown_pdf.markdown_to_pdf(text=MD_SAMPLE, output_path=str(self.work / "bad.pdf"),
                                                     font_size=bad)
                self.assertFalse(result["success"])
                self.assertIn(needle, result["error"])

    def test_bad_page_size(self):
        result = markdown_pdf.markdown_to_pdf(text=MD_SAMPLE, output_path=str(self.work / "bad.pdf"),
                                             page_size="a3")
        self.assertFalse(result["success"])
        self.assertIn("不支持的纸张尺寸", result["error"])

    def test_empty_text(self):
        result = markdown_pdf.markdown_to_pdf(text="   \n\n", output_path=str(self.work / "empty.pdf"))
        self.assertFalse(result["success"])
        self.assertIn("没有可转换的 Markdown 内容", result["error"])

    def test_file_input_uses_file_name(self):
        md_file = self.work / "说明文档.md"
        md_file.write_text("# 来自文件的中文标题\n\n正文内容。\n", encoding="utf-8")
        result = markdown_pdf.markdown_to_pdf(source_file=str(md_file), output_dir=str(self.work))
        self.assertTrue(result["success"], result.get("error"))
        self.assertTrue(result["filename"].endswith(".pdf"))
        doc = fitz.open(result["output"])
        try:
            text = doc[0].get_text()
        finally:
            doc.close()
        self.assertIn("来自文件的中文标题", text)

    def test_file_wins_over_text(self):
        md_file = self.work / "priority.md"
        md_file.write_text("# 文件里的内容优先\n", encoding="utf-8")
        result = markdown_pdf.markdown_to_pdf(text="# 这段文本不该出现\n",
                                             source_file=str(md_file),
                                             output_dir=str(self.work))
        self.assertTrue(result["success"], result.get("error"))
        doc = fitz.open(result["output"])
        try:
            text = doc[0].get_text()
        finally:
            doc.close()
        self.assertIn("文件里的内容优先", text)
        self.assertNotIn("这段文本不该出现", text)

    def test_missing_md_file(self):
        result = markdown_pdf.markdown_to_pdf(source_file=str(self.work / "nope.md"),
                                             output_dir=str(self.work))
        self.assertFalse(result["success"])
        self.assertIn("找不到 Markdown 文件", result["error"])

    def test_gbk_markdown_file(self):
        md_file = self.work / "gbk.md"
        md_file.write_bytes("# 中文标题测试\n\n这是 GBK 编码的正文。\n".encode("gb18030"))
        result = markdown_pdf.markdown_to_pdf(source_file=str(md_file), output_dir=str(self.work))
        self.assertTrue(result["success"], result.get("error"))
        doc = fitz.open(result["output"])
        try:
            text = doc[0].get_text()
        finally:
            doc.close()
        self.assertIn("中文标题测试", text)


# ══════════════════════════════════════════════════════════════════════
# 6) 走真实 tasks.py 分发链路（证明前端能拿到结果）
# ══════════════════════════════════════════════════════════════════════
class NewToolJobPipelineTests(WorkerToolTest):
    def test_video_trim_job(self):
        video = make_video(self.work / "clip.mp4", duration=5, gop=10)
        job = self.run_job("video-trim", {"file": video, "start": "1", "end": "3", "mode": "precise"})
        path = self.assert_job_ok(job)
        self.assertEqual(job["resultFilename"], "clip-trimmed.mp4")
        self.assertEqual(job["resultMimeType"], "video/mp4")
        self.assertAlmostEqual(probe_duration(path), 2.0, delta=0.35)

    def test_video_trim_job_failure_has_chinese_error(self):
        video = make_video(self.work / "clip2.mp4", duration=5)
        job = self.run_job("video-trim", {"file": video, "start": "4", "end": "2"})
        self.assert_job_failed(job, "结束时间必须晚于开始时间")

    def test_audio_trim_job(self):
        audio = make_audio(self.work / "tone.mp3", duration=4)
        job = self.run_job("audio-trim", {"file": audio, "start": "1", "end": "3", "mode": "fast"})
        path = self.assert_job_ok(job)
        self.assertEqual(job["resultFilename"], "tone-trimmed.mp3")
        self.assertEqual(job["resultMimeType"], "audio/mpeg")
        self.assertAlmostEqual(probe_duration(path), 2.0, delta=0.5)

    def test_frame_extract_job(self):
        video = make_video(self.work / "clip3.mp4", duration=5)
        job = self.run_job("video-frame-extract",
                           {"file": video, "time": "1.5", "format": "jpg", "width": "160"})
        path = self.assert_job_ok(job)
        self.assertEqual(job["resultFilename"], "clip3-frame.jpg")
        self.assertEqual(job["resultMimeType"], "image/jpeg")
        with Image.open(path) as img:
            self.assertEqual(img.size, (160, 120))

    def test_frame_extract_job_png_defaults(self):
        video = make_video(self.work / "clip4.mp4", duration=3)
        job = self.run_job("video-frame-extract", {"file": video})
        path = self.assert_job_ok(job)
        self.assertEqual(job["resultFilename"], "clip4-frame.png")
        self.assertEqual(job["resultMimeType"], "image/png")
        with Image.open(path) as img:
            self.assertEqual(img.size, (320, 240))

    def test_csv_to_xlsx_job(self):
        csv = write_csv(self.work / "data.csv", CHINESE_CSV)
        job = self.run_job("csv-excel", {"file": csv, "direction": "auto"})
        path = self.assert_job_ok(job)
        self.assertEqual(job["resultFilename"], "data.xlsx")
        self.assertIn("spreadsheetml", job["resultMimeType"])
        wb = load_workbook(path)
        self.assertEqual(wb.active["A2"].value, "张三")
        wb.close()

    def test_csv_to_xlsx_job_with_has_header_false(self):
        csv = write_csv(self.work / "noheader.csv", "姓名,年龄\n张三,30\n")
        job = self.run_job("csv-excel", {"file": csv, "direction": "to-xlsx", "has_header": "false"})
        path = self.assert_job_ok(job)
        wb = load_workbook(path)
        ws = wb.active
        self.assertFalse(ws["A1"].font.bold, "payload 里的 has_header=false 必须真的传到工具里")
        self.assertIsNone(ws.freeze_panes)
        self.assertEqual(ws["A1"].value, "姓名")
        wb.close()

    def test_csv_to_xlsx_job_with_has_header_true(self):
        csv = write_csv(self.work / "withheader.csv", "姓名,年龄\n张三,30\n")
        job = self.run_job("csv-excel", {"file": csv, "has_header": "true"})
        path = self.assert_job_ok(job)
        wb = load_workbook(path)
        self.assertTrue(wb.active["A1"].font.bold)
        self.assertEqual(wb.active.freeze_panes, "A2")
        wb.close()

    def test_xlsx_to_csv_job(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["姓名", "年龄"])
        ws.append(["张三", 30])
        xlsx = self.work / "book.xlsx"
        wb.save(str(xlsx))
        wb.close()
        job = self.run_job("csv-excel", {"file": str(xlsx), "direction": "auto"})
        path = self.assert_job_ok(job)
        self.assertEqual(job["resultFilename"], "book.csv")
        self.assertEqual(job["resultMimeType"], "text/csv")
        self.assertIn("张三", path.read_text(encoding="utf-8-sig"))

    def test_wrong_sheet_error_is_chinese(self):
        wb = Workbook()
        wb.active.append(["a"])
        xlsx = self.work / "onlyOne.xlsx"
        wb.save(str(xlsx))
        wb.close()
        job = self.run_job("csv-excel", {"file": str(xlsx), "sheet": "不存在"})
        self.assert_job_failed(job, "找不到工作表")

    def test_markdown_to_pdf_job_from_text(self):
        job = self.run_job("markdown-to-pdf",
                           {"text": "# 中文标题\n\n这是**正文**。\n", "page_size": "a4", "font_size": "12"})
        path = self.assert_job_ok(job)
        self.assertEqual(job["resultFilename"], "document.pdf")
        self.assertEqual(job["resultMimeType"], "application/pdf")
        doc = fitz.open(str(path))
        try:
            text = doc[0].get_text()
        finally:
            doc.close()
        self.assertIn("中文标题", text)

    def test_markdown_to_pdf_job_from_file(self):
        md_file = self.work / "readme.md"
        md_file.write_text("# 文件入口的中文标题\n\n正文。\n", encoding="utf-8")
        job = self.run_job("markdown-to-pdf", {"file": str(md_file), "font_size": 14})
        path = self.assert_job_ok(job)
        self.assertEqual(job["resultFilename"], "readme.pdf")
        doc = fitz.open(str(path))
        try:
            text = doc[0].get_text()
        finally:
            doc.close()
        self.assertIn("文件入口的中文标题", text)


if __name__ == "__main__":
    unittest.main(verbosity=2)
