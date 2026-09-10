"""
PDF 转换工具处理模块
参考 docsmall.com 的 PDF 转换工具逻辑
"""
import os
import io
from typing import List, Dict, Any, Optional
import pymupdf as fitz
from PIL import Image


def pdf_to_word(file_path: str, output_path: str) -> Dict[str, Any]:
    """PDF 转 Word"""
    try:
        from pdf2docx import Converter
        cv = Converter(file_path)
        cv.convert(output_path, start=0, end=None)
        cv.close()
        return {"success": True, "output": output_path}
    except ImportError:
        # 如果 pdf2docx 不可用，用备用方案：提取文本写入 Word
        from docx import Document
        doc = fitz.open(file_path)
        word_doc = Document()
        
        for page in doc:
            text = page.get_text()
            if text.strip():
                word_doc.add_paragraph(text)
        
        word_doc.save(output_path)
        doc.close()
        return {"success": True, "output": output_path, "note": "使用文本提取模式（未安装 pdf2docx）"}


def pdf_to_excel(file_path: str, output_path: str) -> Dict[str, Any]:
    """PDF 转 Excel"""
    try:
        import tabula
        # 尝试用 tabula 提取表格
        tables = tabula.read_pdf(file_path, pages="all", multiple_tables=True)
        if tables:
            import pandas as pd
            with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
                for i, table in enumerate(tables):
                    table.to_excel(writer, sheet_name=f"Table_{i+1}", index=False)
            return {"success": True, "output": output_path, "tables": len(tables)}
    except Exception:
        pass
    
    # 备用方案：提取所有文本写入 Excel
    from openpyxl import Workbook
    doc = fitz.open(file_path)
    wb = Workbook()
    ws = wb.active
    ws.title = "PDF内容"
    
    row = 1
    for page_num, page in enumerate(doc, 1):
        text = page.get_text()
        ws.cell(row=row, column=1, value=f"第 {page_num} 页")
        row += 1
        for line in text.split("\n"):
            if line.strip():
                ws.cell(row=row, column=1, value=line)
                row += 1
        row += 1
    
    wb.save(output_path)
    doc.close()
    return {"success": True, "output": output_path, "note": "使用文本提取模式"}


def pdf_to_ppt(file_path: str, output_path: str, dpi: int = 150) -> Dict[str, Any]:
    """PDF 转 PPT（每页转为图片插入）"""
    from pptx import Presentation
    from pptx.util import Inches
    
    doc = fitz.open(file_path)
    prs = Presentation()
    
    # 设置幻灯片大小为 A4
    prs.slide_width = Inches(8.27)
    prs.slide_height = Inches(11.69)
    
    zoom = dpi / 72
    matrix = fitz.Matrix(zoom, zoom)
    
    blank_layout = prs.slide_layouts[6]  # 空白布局
    
    for i, page in enumerate(doc):
        pix = page.get_pixmap(matrix=matrix)
        img_path = os.path.join(os.path.dirname(output_path), f"temp_page_{i+1}.png")
        pix.save(img_path)
        
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(
            img_path,
            Inches(0), Inches(0),
            width=prs.slide_width,
            height=prs.slide_height
        )
        
        # 删除临时图片
        os.remove(img_path)
    
    pages = len(doc)
    prs.save(output_path)
    doc.close()
    return {"success": True, "output": output_path, "pages": pages}


def pdf_to_images(file_path: str, output_dir: str, format: str = "png", dpi: int = 150) -> Dict[str, Any]:
    """PDF 转图片"""
    doc = fitz.open(file_path)
    total_pages = len(doc)
    outputs = []
    
    zoom = dpi / 72
    matrix = fitz.Matrix(zoom, zoom)
    
    for i, page in enumerate(doc):
        pix = page.get_pixmap(matrix=matrix)
        output_path = os.path.join(output_dir, f"page_{i+1:03d}.{format}")
        pix.save(output_path)
        outputs.append(output_path)
    
    doc.close()
    return {"success": True, "outputs": outputs, "count": total_pages, "format": format, "dpi": dpi}


def image_to_word(file_path: str, output_path: str) -> Dict[str, Any]:
    """图片转 Word（将图片插入 Word 文档）"""
    from docx import Document
    from docx.shared import Inches
    
    doc = Document()
    doc.add_picture(file_path, width=Inches(6))
    doc.save(output_path)
    return {"success": True, "output": output_path}


def image_to_excel(file_path: str, output_path: str) -> Dict[str, Any]:
    """图片转 Excel（将图片插入 Excel 表格）"""
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    
    wb = Workbook()
    ws = wb.active
    ws.title = "图片"
    
    img = XLImage(file_path)
    # 按比例缩放
    max_width = 600
    if img.width > max_width:
        ratio = max_width / img.width
        img.width = max_width
        img.height = int(img.height * ratio)
    
    ws.add_image(img, "A1")
    wb.save(output_path)
    return {"success": True, "output": output_path}
