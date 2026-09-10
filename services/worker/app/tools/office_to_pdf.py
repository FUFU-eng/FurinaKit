"""
Office 文档转 PDF 处理模块
支持 Word、Excel、PPT 转 PDF
"""
import os
import tempfile
import logging
from typing import Dict, Any

logger = logging.getLogger("furinakit.office")


def word_to_pdf(file_path: str, output_path: str) -> Dict[str, Any]:
    """Word 转 PDF"""
    try:
        # 尝试使用 docx2pdf（需要 Microsoft Word）
        from docx2pdf import convert
        convert(file_path, output_path)
        return {"success": True, "output": output_path, "method": "docx2pdf"}
    except Exception as e:
        logger.warning("docx2pdf 失败: %s", e)
    
    # 备用方案：使用 LibreOffice
    try:
        import subprocess
        result = subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", "--outdir", os.path.dirname(output_path), file_path],
            capture_output=True, text=True, timeout=60,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
        if result.returncode == 0:
            # 重命名输出文件
            generated = os.path.join(os.path.dirname(output_path), os.path.splitext(os.path.basename(file_path))[0] + ".pdf")
            if os.path.exists(generated):
                os.rename(generated, output_path)
            return {"success": True, "output": output_path, "method": "libreoffice"}
    except Exception as e:
        logger.warning("LibreOffice 失败: %s", e)
    
    # 最后备用方案：提取文本生成简单 PDF
    try:
        import pymupdf as fitz
        from docx import Document
        
        doc = Document(file_path)
        pdf_doc = fitz.open()
        page = pdf_doc.new_page()
        text = "\n".join([para.text for para in doc.paragraphs])
        page.insert_text((50, 50), text, fontsize=12)
        pdf_doc.save(output_path)
        pdf_doc.close()
        return {"success": True, "output": output_path, "method": "text-extraction", "note": "仅提取文本，格式可能丢失"}
    except Exception as e:
        return {"success": False, "error": str(e)}


def excel_to_pdf(file_path: str, output_path: str) -> Dict[str, Any]:
    """Excel 转 PDF"""
    try:
        # 尝试使用 win32com（需要 Microsoft Excel）
        import win32com.client
        excel = win32com.client.Dispatch("Excel.Application")
        excel.Visible = False
        wb = excel.Workbooks.Open(os.path.abspath(file_path))
        wb.ExportAsFixedFormat(0, os.path.abspath(output_path))
        wb.Close()
        excel.Quit()
        return {"success": True, "output": output_path, "method": "excel-com"}
    except Exception as e:
        logger.warning("Excel COM 失败: %s", e)
    
    # 备用方案：使用 LibreOffice
    try:
        import subprocess
        result = subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", "--outdir", os.path.dirname(output_path), file_path],
            capture_output=True, text=True, timeout=60,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
        if result.returncode == 0:
            generated = os.path.join(os.path.dirname(output_path), os.path.splitext(os.path.basename(file_path))[0] + ".pdf")
            if os.path.exists(generated):
                os.rename(generated, output_path)
            return {"success": True, "output": output_path, "method": "libreoffice"}
    except Exception as e:
        logger.warning("LibreOffice 失败: %s", e)
    
    # 最后备用方案：提取数据生成简单 PDF
    try:
        import pymupdf as fitz
        from openpyxl import load_workbook
        
        wb = load_workbook(file_path, data_only=True)
        pdf_doc = fitz.open()
        
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            page = pdf_doc.new_page()
            y = 50
            page.insert_text((50, y), f"Sheet: {sheet_name}", fontsize=14)
            y += 25
            
            for row in ws.iter_rows(values_only=True):
                line = " | ".join([str(cell) if cell is not None else "" for cell in row])
                if line.strip(" |"):
                    page.insert_text((50, y), line[:100], fontsize=10)
                    y += 15
                    if y > 800:
                        page = pdf_doc.new_page()
                        y = 50
        
        pdf_doc.save(output_path)
        pdf_doc.close()
        return {"success": True, "output": output_path, "method": "data-extraction", "note": "仅提取数据，格式可能丢失"}
    except Exception as e:
        return {"success": False, "error": str(e)}


def ppt_to_pdf(file_path: str, output_path: str) -> Dict[str, Any]:
    """PPT 转 PDF"""
    try:
        # 尝试使用 win32com（需要 Microsoft PowerPoint）
        import win32com.client
        powerpoint = win32com.client.Dispatch("PowerPoint.Application")
        presentation = powerpoint.Presentations.Open(os.path.abspath(file_path), WithWindow=False)
        presentation.SaveAs(os.path.abspath(output_path), 32)  # 32 = ppSaveAsPDF
        presentation.Close()
        powerpoint.Quit()
        return {"success": True, "output": output_path, "method": "powerpoint-com"}
    except Exception as e:
        logger.warning("PowerPoint COM 失败: %s", e)
    
    # 备用方案：使用 LibreOffice
    try:
        import subprocess
        result = subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", "--outdir", os.path.dirname(output_path), file_path],
            capture_output=True, text=True, timeout=60,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
        if result.returncode == 0:
            generated = os.path.join(os.path.dirname(output_path), os.path.splitext(os.path.basename(file_path))[0] + ".pdf")
            if os.path.exists(generated):
                os.rename(generated, output_path)
            return {"success": True, "output": output_path, "method": "libreoffice"}
    except Exception as e:
        logger.warning("LibreOffice 失败: %s", e)
    
    # 最后备用方案：提取文本生成简单 PDF
    try:
        import pymupdf as fitz
        from pptx import Presentation
        
        prs = Presentation(file_path)
        pdf_doc = fitz.open()
        
        for i, slide in enumerate(prs.slides):
            page = pdf_doc.new_page()
            y = 50
            page.insert_text((50, y), f"Slide {i+1}", fontsize=16)
            y += 30
            
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    for line in shape.text.split("\n"):
                        if line.strip():
                            page.insert_text((50, y), line[:80], fontsize=11)
                            y += 18
                            if y > 800:
                                page = pdf_doc.new_page()
                                y = 50
        
        pdf_doc.save(output_path)
        pdf_doc.close()
        return {"success": True, "output": output_path, "method": "text-extraction", "note": "仅提取文本，格式可能丢失"}
    except Exception as e:
        return {"success": False, "error": str(e)}
