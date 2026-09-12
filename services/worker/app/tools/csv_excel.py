"""CSV ↔ Excel 双向转换（tool id: csv-excel），真正写出 .xlsx（openpyxl）。

重点：
- CSV → xlsx：自动探测分隔符、兼容中文 CSV 编码（UTF-8/BOM/GBK/GB18030/UTF-16）、
  首行作表头（加粗 + 冻结首行 + 按内容估算列宽，可用 has_header=false 关掉）、
  数字写成**真数字**，但「看起来像编号」的值（前导 0、超过 15 位的长数字）保持文本，
  避免 Excel 破坏精度
- xlsx → CSV：可选工作表，输出 UTF-8 带 BOM（Excel 打开不乱码），正确转义逗号/引号/换行
- 大文件保护：CSV 体积 / 行数上限提示，xlsx 逐行流式读取，不把整个表读进内存
- 本文件是新增模块，不改动任何现有模块的行为
"""

import csv
import io
import re
from collections import Counter
from datetime import date, datetime, time as _time, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# 上限保护（超出给中文提示，而不是把内存吃爆 / 卡死）
MAX_CSV_BYTES = 100 * 1024 * 1024      # 100 MB
MAX_CSV_ROWS = 100_000                 # 10 万行
MAX_XLSX_ROWS = 1_000_000              # Excel 自身行上限 1048576
MAX_COLUMNS = 16384                    # Excel 列上限
MIN_COL_WIDTH = 8.0
MAX_COL_WIDTH = 60.0

CANDIDATE_DELIMITERS = (",", ";", "\t", "|")
_DELIMITER_LABELS = {",": "逗号", ";": "分号", "\t": "Tab", "|": "竖线", " ": "空格"}

# has_header 的容错取值（大小写不敏感）；非法值一律按 True（有表头）处理，不为这一个开关报错
_TRUE_VALUES = {"true", "1", "yes", "y", "on", "t", "是", "有", "表头"}
_FALSE_VALUES = {"false", "0", "no", "n", "off", "f", "否", "无", "无表头"}

_FLOAT_RE = re.compile(r"^[+-]?(?:\d+\.\d+|\.\d+|\d+)(?:[eE][+-]?\d+)?$")

# 各种 BOM：Excel 导出的 CSV 常见 utf-8-sig / utf-16
_BOMS = (
    (b"\xef\xbb\xbf", "utf-8-sig"),
    (b"\xff\xfe\x00\x00", "utf-32"),
    (b"\x00\x00\xfe\xff", "utf-32"),
    (b"\xff\xfe", "utf-16"),
    (b"\xfe\xff", "utf-16"),
)


class CsvExcelError(Exception):
    """可以直接展示给用户的中文错误。"""


def failed(exc: BaseException, prefix: str = "") -> Dict[str, Any]:
    if isinstance(exc, CsvExcelError):
        return {"success": False, "error": str(exc)}
    return {"success": False, "error": f"{prefix}：{exc}" if prefix else f"处理失败：{exc}"}


# ── 编码 / 分隔符 ────────────────────────────────────────────────────

def decode_csv_bytes(data: bytes) -> Tuple[str, str]:
    """按 UTF-8（含 BOM）→ GB18030 的顺序解码，返回 (文本, 使用的编码)。"""
    for bom, enc in _BOMS:
        if data.startswith(bom):
            try:
                return data.decode(enc), enc
            except UnicodeDecodeError:
                break
    try:
        return data.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        pass
    try:
        return data.decode("gb18030"), "gb18030（兼容 GBK）"
    except UnicodeDecodeError:
        pass
    return data.decode("utf-8", errors="replace"), "utf-8（部分字符无法识别）"


def detect_delimiter(text: str, sample_lines: int = 20) -> str:
    """看前几行里哪种分隔符切出来的列数最整齐（一致率最高，其次列数最多）。"""
    lines = [ln for ln in text.splitlines() if ln.strip()][:sample_lines]
    if not lines:
        return ","
    sample = "\n".join(lines)
    best: Tuple[Optional[str], float, int] = (None, 0.0, 0)
    for delim in CANDIDATE_DELIMITERS:
        try:
            rows = [r for r in csv.reader(io.StringIO(sample), delimiter=delim) if r]
        except Exception:
            continue
        counts = [len(r) for r in rows]
        if not counts:
            continue
        top_count, top_freq = Counter(counts).most_common(1)[0]
        if top_count < 2:
            continue
        consistency = top_freq / len(counts)
        if best[0] is None or (consistency, top_count) > (best[1], best[2]):
            best = (delim, consistency, top_count)
    return best[0] or ","


def normalize_delimiter(value: Any) -> Optional[str]:
    """payload 里的 delimiter：空 / auto 表示自动探测。"""
    if value is None:
        return None
    raw = str(value)
    lowered = raw.strip().lower()
    if lowered in ("", "auto", "自动", "自动探测"):
        return None
    named = {
        "tab": "\t", "\\t": "\t", "制表符": "\t",
        "comma": ",", "逗号": ",",
        "semicolon": ";", "分号": ";",
        "pipe": "|", "竖线": "|",
        "space": " ", "空格": " ",
    }
    if lowered in named:
        return named[lowered]
    if len(raw) == 1:
        return raw
    raise CsvExcelError(
        f"分隔符只能是单个字符（例如 , ; 或 Tab），收到的是「{value}」"
    )


def parse_has_header(value: Any) -> bool:
    """解析 payload 里的 has_header（首行是否为表头）。

    缺省 / 空 / 非法值一律当作 True（有表头）；只为「明确表示否」的值返回 False。
    契约要求：绝不因为这一个开关报错。
    """
    if value is None:
        return True
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    raw = str(value).strip().lower()
    if raw == "":
        return True
    if raw in _FALSE_VALUES:
        return False
    return True


def resolve_direction(path: Any, direction: Any = "auto") -> str:
    """返回 "to-xlsx" 或 "to-csv"。"""
    raw = str(direction if direction is not None else "auto").strip().lower()
    ext = Path(str(path or "")).suffix.lower()
    if raw in ("", "auto", "自动"):
        if ext == ".csv":
            return "to-xlsx"
        if ext in (".xlsx", ".xlsm"):
            return "to-csv"
        raise CsvExcelError(
            f"无法根据扩展名判断转换方向：{ext or '（没有扩展名）'}，"
            "请明确指定 direction=to-xlsx 或 to-csv"
        )
    if raw in ("to-xlsx", "to_excel"):
        return "to-xlsx"
    if raw in ("to-csv", "to_csv"):
        return "to-csv"
    raise CsvExcelError(
        f"无法识别的转换方向：「{direction}」，只支持 auto / to-xlsx / to-csv"
    )


# ── 单元格取值 ──────────────────────────────────────────────────────

def to_cell_value(raw: Optional[str]) -> Any:
    """把 CSV 文本转成写入 xlsx 的值。

    - 空字符串 → None（真空单元格）
    - 普通数字 → int / float（这样在 Excel 里能直接求和）
    - 前导 0 的编号、超过 15 位的长数字 → 原样保留为文本（否则 Excel 会丢精度）
    - 其它一律文本，且保留原始内容（不裁剪空格）
    """
    if raw is None:
        return None
    text = str(raw)
    stripped = text.strip()
    if stripped == "":
        return None

    body = stripped[1:] if stripped[:1] in "+-" else stripped
    if not body:
        return text

    # 1) 以 0 开头的编号（007、00123）必须保持文本
    if len(body) > 1 and body[0] == "0" and body[1].isdigit():
        return text
    # 2) 纯整数：超过 15 位有效数字 Excel 会丢精度 → 文本
    if body.isascii() and body.isdigit():
        if len(body) > 15:
            return text
        try:
            return int(stripped)
        except ValueError:
            return text
    # 3) 小数 / 科学计数法
    if _FLOAT_RE.match(stripped):
        mantissa = re.split(r"[eE]", body)[0].replace(".", "")
        if len(mantissa) > 15:
            return text
        try:
            value = float(stripped)
        except ValueError:
            return text
        if value != value or value in (float("inf"), float("-inf")):
            return text
        return value
    return text


def _display_width(text: str) -> int:
    """估算列宽：中日韩字符按 2 个字符宽算。"""
    width = 0
    for ch in text:
        width += 2 if ("\u2e80" <= ch <= "\u9fff" or "\uf900" <= ch <= "\ufaff" or "\uff00" <= ch <= "\uffef") else 1
    return width


def _safe_sheet_title(title: Optional[str]) -> Optional[str]:
    if title is None:
        return None
    raw = str(title).strip()
    if raw == "":
        return None
    cleaned = re.sub(r"[\[\]:*?/\\]", "_", raw)[:31].strip()
    return cleaned or None


# ── CSV → xlsx ──────────────────────────────────────────────────────

def csv_to_xlsx(
    csv_path: Any,
    xlsx_path: Any,
    sheet: Optional[str] = None,
    delimiter: Any = None,
    has_header: bool = True,
) -> Dict[str, Any]:
    src = Path(csv_path)
    try:
        size = src.stat().st_size
    except OSError as exc:
        raise CsvExcelError(f"无法读取 CSV 文件：{src.name}（{exc}）") from exc
    if size > MAX_CSV_BYTES:
        raise CsvExcelError(
            f"CSV 文件太大（{size / 1048576:.1f} MB），超过 {MAX_CSV_BYTES // 1048576} MB 上限，"
            "请先拆分后再转换"
        )

    try:
        data = src.read_bytes()
    except OSError as exc:
        raise CsvExcelError(f"无法读取 CSV 文件：{src.name}（{exc}）") from exc
    if not data.strip():
        raise CsvExcelError("CSV 文件是空的，没有可转换的内容")

    text, encoding = decode_csv_bytes(data)
    explicit = normalize_delimiter(delimiter)
    delim = explicit if explicit is not None else detect_delimiter(text)

    rows: List[List[str]] = []
    try:
        reader = csv.reader(io.StringIO(text), delimiter=delim)
        for row in reader:
            rows.append(["" if c is None else str(c) for c in row])
            if len(rows) > MAX_CSV_ROWS:
                raise CsvExcelError(
                    f"CSV 行数太多（超过 {MAX_CSV_ROWS:,} 行上限），请先拆分后再转换"
                )
    except csv.Error as exc:
        raise CsvExcelError(f"CSV 解析失败（第 {reader.line_num} 行附近）：{exc}") from exc

    while rows and all(c.strip() == "" for c in rows[-1]):
        rows.pop()
    if not rows:
        raise CsvExcelError("CSV 文件里没有有效内容")

    col_count = max(len(r) for r in rows)
    if col_count > MAX_COLUMNS:
        raise CsvExcelError(f"CSV 列数太多（{col_count} 列），超过 Excel 上限 {MAX_COLUMNS} 列")

    wb = Workbook()
    ws = wb.active
    ws.title = _safe_sheet_title(sheet) or "Sheet1"

    widths = [0] * col_count
    first_data_row = 1
    if has_header and rows:
        # 首行作表头：加粗 + 冻结首行 + 用表头内容参与列宽估算
        header = rows[0]
        for col in range(col_count):
            raw = header[col] if col < len(header) else ""
            label = _truncate_header(raw.strip() or f"列{col + 1}")
            cell = ws.cell(row=1, column=col + 1)
            _assign(cell, label)
            cell.font = Font(bold=True)
            widths[col] = max(widths[col], _display_width(label))
        first_data_row = 2

    data_rows = rows[1:] if has_header else rows
    for offset, row in enumerate(data_rows):
        r_index = first_data_row + offset
        for col in range(col_count):
            raw = row[col] if col < len(row) else ""
            value = to_cell_value(raw)
            if value is None:
                continue
            cell = ws.cell(row=r_index, column=col + 1)
            _assign(cell, value)
            widths[col] = max(widths[col], _display_width(str(value)))

    if has_header:
        ws.freeze_panes = "A2"
    for col in range(col_count):
        ws.column_dimensions[get_column_letter(col + 1)].width = min(
            MAX_COL_WIDTH, max(MIN_COL_WIDTH, widths[col] + 2)
        )

    out = Path(xlsx_path)
    if out.parent and str(out.parent):
        out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    wb.close()

    label = _DELIMITER_LABELS.get(delim, repr(delim))
    return {
        "success": True,
        "output": str(out),
        "filename": out.name,
        "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "message": (
            f"已转换为 Excel（{len(data_rows)} 行数据 × {col_count} 列；"
            f"{'首行作表头' if has_header else '首行按普通数据'}"
            f"；识别编码 {encoding}，分隔符 {label}）"
        ),
    }


def _truncate_header(label: str, limit: int = 200) -> str:
    return label if len(label) <= limit else label[:limit] + "…"


def _assign(cell, value: Any) -> None:
    """写单元格；以 = 开头的内容强制按文本写入，避免 CSV 注入被 Excel 当成公式执行。"""
    cell.value = value
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"


# ── xlsx → CSV ──────────────────────────────────────────────────────

def _cell_to_csv_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return repr(value)
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, (_time, date)):
        return value.isoformat()
    if isinstance(value, timedelta):
        total = int(value.total_seconds())
        sign = "-" if total < 0 else ""
        total = abs(total)
        return f"{sign}{total // 3600:02d}:{total % 3600 // 60:02d}:{total % 60:02d}"
    if isinstance(value, Decimal):
        return format(value, "f")
    return str(value)


def xlsx_to_csv(xlsx_path: Any, csv_path: Any, sheet: Optional[str] = None,
                has_header: bool = True) -> Dict[str, Any]:
    """把工作表导出成 UTF-8 带 BOM 的 CSV。

    关于 has_header：导出方向本来就不会对首行做任何特殊处理（始终按单元格原样写出），
    所以 True / False 的产物逐字节一致；这里保留参数只是为了和 CSV→xlsx 的契约对称，
    并在成功消息里如实说明首行是怎么被理解的。
    """
    src = Path(xlsx_path)
    try:
        wb = load_workbook(str(src), read_only=True, data_only=True, keep_links=False)
    except Exception as exc:
        raise CsvExcelError(
            f"无法读取 Excel 文件：{src.name}（请确认它是 .xlsx / .xlsm 格式，"
            f"老的 .xls 需要先另存为 .xlsx）"
        ) from exc

    try:
        names = list(wb.sheetnames)
        if not names:
            raise CsvExcelError("这个 Excel 文件里没有任何工作表")

        wanted = str(sheet).strip() if sheet is not None else ""
        if wanted:
            if wanted in names:
                ws = wb[wanted]
            elif _safe_sheet_title(wanted) in names:
                ws = wb[_safe_sheet_title(wanted)]
            else:
                raise CsvExcelError(
                    f"找不到工作表「{sheet}」，这个文件里的工作表有：{'、'.join(names)}"
                )
        else:
            ws = wb[names[0]]
        used_sheet = ws.title

        out = Path(csv_path)
        if out.parent and str(out.parent):
            out.parent.mkdir(parents=True, exist_ok=True)

        row_count = 0
        with open(out, "w", encoding="utf-8-sig", newline="") as fh:
            writer = csv.writer(fh, lineterminator="\r\n")
            for row in ws.iter_rows(values_only=True):
                row_count += 1
                if row_count > MAX_XLSX_ROWS:
                    raise CsvExcelError(
                        f"工作表「{used_sheet}」行数超过 {MAX_XLSX_ROWS:,} 行，无法导出为 CSV"
                    )
                cells = [_cell_to_csv_text(v) for v in row]
                while cells and cells[-1] == "":
                    cells.pop()
                writer.writerow(cells)
        if row_count == 0:
            raise CsvExcelError(f"工作表「{used_sheet}」里没有任何内容")
    finally:
        try:
            wb.close()
        except Exception:
            pass

    return {
        "success": True,
        "output": str(out),
        "filename": out.name,
        "mime": "text/csv",
        "message": (
            f"已转换为 CSV（工作表「{used_sheet}」，{row_count} 行，"
            f"{'首行作表头' if has_header else '首行按普通数据'}，UTF-8 带 BOM，Excel 打开不乱码）"
        ),
    }


# ── 统一入口 ────────────────────────────────────────────────────────

def convert_file(
    input_path: Any,
    output_dir: Any,
    direction: Any = "auto",
    sheet: Optional[str] = None,
    delimiter: Any = None,
    has_header: Any = None,
) -> Dict[str, Any]:
    """csv-excel 的统一入口：按 direction / 扩展名决定转换方向。

    has_header：首行是否为表头（"true"/"false"，缺省 true，非法值按 true）。
    """
    try:
        raw = str(input_path or "").strip()
        if not raw:
            raise CsvExcelError("没有拿到要转换的文件路径，请重新选择文件")
        src = Path(raw)
        if not src.is_file():
            raise CsvExcelError(f"找不到文件：{raw}")

        resolved = resolve_direction(src, direction)
        out_dir = Path(str(output_dir or "."))
        out_dir.mkdir(parents=True, exist_ok=True)
        header_flag = parse_has_header(has_header)

        if resolved == "to-xlsx":
            return csv_to_xlsx(src, out_dir / f"{src.stem}.xlsx", sheet=sheet,
                               delimiter=delimiter, has_header=header_flag)
        return xlsx_to_csv(src, out_dir / f"{src.stem}.csv", sheet=sheet, has_header=header_flag)
    except CsvExcelError as exc:
        return failed(exc)
    except Exception as exc:  # noqa: BLE001
        return failed(exc, "CSV / Excel 转换失败")
